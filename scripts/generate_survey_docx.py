import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level > 1 else 20)
    p.paragraph_format.space_after = Pt(6)
    run = p.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(24, 43, 73) # Navy
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(41, 74, 110) # Steel Navy
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(60, 60, 60)
    return p

def main():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(5)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("KLAUSE: LITERATURE SURVEY & GAP ANALYSIS")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(24, 43, 73)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(20)
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Foundational Research Papers, Architectural Mapping, and Resolution of Practical Gaps in Autonomous Desktop Engineering Assistants")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(90, 105, 120)

    # Meta banner
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run("Target Project: KLAUSE Desktop Engineering Assistant | Scope: 30 Research Papers | Status: Fully Verified")
    meta_run.font.size = Pt(9.5)
    meta_run.font.bold = True
    meta_run.font.color.rgb = RGBColor(70, 80, 95)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(16)
    p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="182B49"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

    # 1. Executive Summary
    add_header_styled(doc, "1. Executive Summary", level=1)
    doc.add_paragraph(
        "KLAUSE is a personal, autonomous AI engineering companion and desktop assistant for Windows. "
        "It uniquely bridges high-level cognitive agent architectures with local desktop execution, "
        "combining an asynchronous ReAct brain, persistent hierarchical memory (MongoDB and ChromaDB vector stores), "
        "bidirectional personal knowledge graph synchronization via Obsidian vaults, OS-level window and process automation, "
        "screen perception via OCR, and local push-to-talk voice activation."
    )
    doc.add_paragraph(
        "This document conducts a rigorous academic literature survey of 30 landmark research papers across 8 core pillars "
        "of autonomous agent research. It reviews the theoretical methodologies of each work, relates them directly to "
        "subsystems within the KLAUSE codebase, and presents an exhaustive comparative gap analysis identifying limitations in "
        "the state-of-the-art and how KLAUSE concretely resolves them."
    )

    # 2. Review of 30 Foundational Papers
    add_header_styled(doc, "2. Comprehensive Review of 30 Foundational Papers", level=1)

    pillars = [
        ("Pillar 1: Core Agentic Reasoning, ReAct & Planning Loops", [
            ("1. ReAct: Synergizing Reasoning and Acting in Language Models",
             "Yao, S., Zhao, J., Yu, D., Du, N., Shafran, I., Narasimhan, K., & Cao, Y. (2023)",
             "ICLR 2023 | arXiv:2210.03629 | Status: Peer-Reviewed",
             "app/core/brain.py, app/core/dispatcher.py",
             "Establishes an interleaved loop where language models generate verbal reasoning traces (Thought) alongside task-specific actions (Action) and receive environment feedback (Observation). This dynamic synergy significantly curbs hallucination rates, promotes error recovery, and forms the core cognitive cycle of KLAUSE."),

            ("2. Reflexion: Language Agents with Verbal Reinforcement Learning",
             "Shinn, N., Cassano, F., Berman, E., Gopinath, A., Narasimhan, K., & Yao, S. (2023)",
             "NeurIPS 2023 | arXiv:2303.11366 | Status: Peer-Reviewed",
             "app/core/dispatcher.py, app/vision/error_reader.py",
             "Introduces verbal self-reflection as an alternative to parameter updates. When an agent experiences execution failure, it writes a diagnostic critique to episodic memory to steer subsequent attempts. KLAUSE uses this paradigm to recover from failed CLI commands and compiler exceptions."),

            ("3. Tree of Thoughts: Deliberate Problem Solving with Large Language Models",
             "Yao, S., Yu, D., Zhao, J., Shafran, I., Griffiths, T. L., Cao, Y., & Narasimhan, K. (2023)",
             "NeurIPS 2023 | arXiv:2305.10601 | Status: Peer-Reviewed",
             "KLAUSE_Build_Plan.md (Multi-phase Planner)",
             "Generalizes chain-of-thought into tree-based search structures (BFS/DFS) with lookahead and backtracking capabilities. Used in KLAUSE to structure multi-step project decompositions and evaluate intermediate tool execution states."),

            ("4. CoALA: A Conceptual Architecture for Language Agents",
             "Sumers, T., Yao, S., Narasimhan, K., & Griffiths, T. L. (2024)",
             "TMLR 2024 (Transactions on Machine Learning Research) | arXiv:2309.02427 | Status: Peer-Reviewed",
             "System Architecture (app/core/, app/memory/, app/tools/)",
             "Formalizes a standardized cognitive architecture separating working memory, long-term episodic/semantic stores, an external action space, and structured decision cycles. Serves as KLAUSE's macro architectural design pattern.")
        ]),

        ("Pillar 2: Tool Use, Function Calling & Environment Grounding", [
            ("5. Toolformer: Language Models Can Teach Themselves to Use Tools",
             "Schick, T., Dwivedi-Yu, J., Dessì, R., Raileanu, R., Lomeli, M., et al. (2023)",
             "NeurIPS 2023 | arXiv:2302.04761 | Status: Peer-Reviewed",
             "app/tools/base.py",
             "Demonstrates self-supervised fine-tuning where an LLM learns to insert API calls into continuous text sequences. Establishes the foundations for KLAUSE's declarative tool interfaces and argument validation schemas."),

            ("6. Gorilla: Large Language Model Connected with Massive APIs",
             "Patil, S. G., Zhang, T., Wang, X., & Gonzalez, J. E. (2023)",
             "arXiv:2305.15334 | Status: Preprint",
             "app/core/tool_router.py",
             "Addresses API hallucination and schema drift across large toolsets using retriever-aware instruction tuning. Informs KLAUSE's tool routing to prevent invalid parameter generation for Windows system utilities."),

            ("7. ToolLLM: Facilitating Large Language Models to Master 16000+ Real-world APIs",
             "Qin, Y., Liang, S., Ye, Y., Zhu, K., Hu, S., Chan, C. M., et al. (2024)",
             "ICLR 2024 | arXiv:2307.16789 | Status: Peer-Reviewed",
             "app/core/skill_repo.py, app/core/tool_router.py",
             "Introduces ToolBench and depth-first search decision trees (DFSDT) for navigating complex multi-branch tool executions. Directly guides KLAUSE's modular skill repository and tool chaining.")
        ]),

        ("Pillar 3: Memory Systems & Vector Retrieval (RAG)", [
            ("8. MemGPT: Towards LLMs as Operating Systems",
             "Packer, C., Wooders, S., Lin, K., Fang, V., Patil, S. G., & Gonzalez, J. E. (2023)",
             "arXiv:2310.08560 | Status: Preprint",
             "app/memory/memory_manager.py, app/memory/database.py",
             "Introduces virtual context management inspired by OS memory paging, moving data between active working context, session recall DB, and archival vector stores via autonomous interrupts. Forms the blueprint for KLAUSE's multi-tiered persistence layer."),

            ("9. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks",
             "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., et al. (2020)",
             "NeurIPS 2020 | arXiv:2005.11401 | Status: Peer-Reviewed",
             "app/memory/chroma_store.py",
             "The foundational RAG formulation combining dense vector retrieval with neural generative models, implemented in KLAUSE via ChromaDB for semantic search across project codebases and technical documentation."),

            ("10. Generative Agents: Interactive Simulacra of Human Behavior",
             "Park, J. S., O'Brien, J. C., Cai, C. J., Morris, M. R., Liang, P., & Bernstein, M. S. (2023)",
             "ACM UIST 2023 | arXiv:2304.03442 | Status: Peer-Reviewed",
             "app/core/context.py",
             "Establishes an associative memory stream scoring experiences based on recency, importance, and query relevance, combined with periodic reflection passes. Guides KLAUSE's conversational history ranking."),

            ("11. Self-RAG: Learning to Retrieve, Generate, and Critique through Self-Reflection",
             "Asai, A., Wu, Z., Wang, Y., Sil, A., & Hajishirzi, H. (2024)",
             "ICLR 2024 | arXiv:2310.11511 | Status: Peer-Reviewed",
             "app/memory/knowledge_base.py",
             "Trains language models to output reflection tokens determining when retrieval is necessary and critiquing the factual support of generated text, implemented in KLAUSE to eliminate irrelevant vector store injections.")
        ]),

        ("Pillar 4: Knowledge Graphs & Obsidian Note Synchronization", [
            ("12. From Local to Global: A Graph RAG Approach to Query-Focused Summarization (GraphRAG)",
             "Edge, D., Trinh, H., Cheng, N., Bradley, J., Chao, A., et al. (2024)",
             "arXiv:2404.16130 (Microsoft Research) | Status: Preprint",
             "app/memory/obsidian_connector.py, KLAUSE_Connections.canvas",
             "Extracts knowledge graphs from text and clusters entities into hierarchical communities (Leiden algorithm) for multi-level corpus summarization. Theoretical basis for KLAUSE's bidirectional Obsidian vault parsing and Canvas flow generation."),

            ("13. HippoRAG: Neurobiologically Inspired Long-Term Memory for Large Language Models",
             "Jiménez Gutiérrez, B., Shu, Y., Gu, Y., Yasunaga, M., & Su, Y. (2024)",
             "NeurIPS 2024 | arXiv:2405.14831 | Status: Peer-Reviewed",
             "app/memory/obsidian_connector.py",
             "Integrates knowledge graphs with Personalized PageRank (mimicking hippocampal indexing) to enable multi-hop associative recall across unlinked notes that semantic vector distance misses."),

            ("14. MindMap: Knowledge Graph Prompting Sparks Graph of Thoughts in Large Language Models",
             "Wen, Y., Wang, Z., & Sun, J. (2024)",
             "ACL 2024 | arXiv:2308.09729 | Status: Peer-Reviewed",
             "app/tools/obsidian_tools.py",
             "Shows that prompting LLMs with structured knowledge subgraphs induces graph-of-thought exploration, validating that graph representations enhance relational reasoning over flat markdown chunks."),

            ("15. KG-Agent: An Efficient Autonomous Agent Framework for Complex Reasoning over Knowledge Graph",
             "Jiang, J., Zhou, K., Zhao, W. X., & Wen, J. R. (2025)",
             "ACL 2025 | arXiv:2402.11163 | Status: Peer-Reviewed",
             "app/tools/kb_tools.py",
             "Formulates multi-hop reasoning over knowledge graphs as programmatic tool execution and dynamic memory updates, guiding KLAUSE's knowledge base querying.")
        ]),

        ("Pillar 5: Windows OS Automation, Desktop Grounding & Computer Use", [
            ("16. UFO: A UI-Focused Agent for Windows OS Interaction",
             "Zhang, C., Li, L., He, S., Zhang, X., Qiao, B., et al. (2024)",
             "arXiv:2402.07939 (Microsoft Research) | Status: Preprint",
             "app/tools/window_tools.py, app/tools/system_tools.py, app/automation/",
             "A dual-agent framework (HostAgent and AppAgent) using Windows UI Automation (UIA) and vision models to navigate and execute actions across native Windows applications. Direct architectural counterpart to KLAUSE's desktop automation layer."),

            ("17. Windows Agent Arena: Evaluating Multi-modal OS Agents at Scale",
             "Bonatti, R., Zhao, D., Bonacci, F., et al. (2024)",
             "arXiv:2409.08264 (Microsoft Research) | Status: Preprint",
             "app/automation/, app/tools/app_tools.py",
             "Benchmarking environment on native Windows 11 evaluating multimodal agents across 150+ realistic tasks across File Explorer, Office, browsers, and developer environments."),

            ("18. OSWorld: Benchmarking Multimodal Agents for Open-Ended Tasks in Real Computer Environments",
             "Xie, T., Zhang, D., Chen, J., Li, X., Zhao, S., et al. (2024)",
             "NeurIPS 2024 | arXiv:2404.07972 | Status: Peer-Reviewed",
             "app/automation/event_engine.py",
             "Interactive computer benchmark highlighting that vision-only agents achieve only ~12% success due to brittle visual grounding, motivating KLAUSE's hybrid API-first + vision-fallback architecture."),

            ("19. Towards General Computer Control: A Multimodal Agent for Red Dead Redemption II as a Case Study (Cradle)",
             "Tan, W., Zhang, W., et al. (2024)",
             "arXiv:2403.03186 | Status: Preprint",
             "app/tools/automation_tools.py",
             "Investigates general computer interaction without software APIs through raw screen perception and mouse/keyboard event synthesis.")
        ]),

        ("Pillar 6: Autonomous Software Engineering & Developer Companions", [
            ("20. SWE-bench: Can Language Models Resolve Real-World GitHub Issues?",
             "Jimenez, C. E., Yang, J., Wettig, A., Yao, S., Pei, K., Press, O., & Narasimhan, K. (2024)",
             "ICLR 2024 | arXiv:2310.06770 | Status: Peer-Reviewed",
             "app/tools/git_tools.py, app/agents/github_client.py",
             "Evaluates LLMs on resolving 2,294 real GitHub issues by traversing multi-file repositories, localizing bugs, writing patches, and verifying fixes with unit tests."),

            ("21. SWE-agent: Agent-Computer Interfaces Enable Automated Software Engineering",
             "Yang, J., Jimenez, C. E., Wettig, A., Lieret, K., Yao, S., Narasimhan, K., & Press, O. (2024)",
             "NeurIPS 2024 | arXiv:2405.15793 | Status: Peer-Reviewed",
             "app/tools/vscode_tools.py, app/tools/terminal_tools.py",
             "Introduces Agent-Computer Interfaces (ACI). Demonstrates that specialized, compact developer commands (windowed file viewing, surgical block replacement) drastically outperform standard bash terminals for LLMs."),

            ("22. MetaGPT: Meta-Programming for A Multi-Agent Collaborative Framework",
             "Hong, S., Zhuge, M., Chen, J., Zheng, X., Hao, C., et al. (2024)",
             "ICLR 2024 | arXiv:2308.00352 | Status: Peer-Reviewed",
             "app/automation/rule_engine.py, rules.md",
             "Incorporates human Standard Operating Procedures (SOPs) into multi-agent workflows to eliminate cascading errors. Inspired KLAUSE's deterministic rules engine."),

            ("23. ChatDev: Communicative Agents for Software Development",
             "Qian, C., Liu, W., Liu, H., Chen, N., Dang, Y., et al. (2024)",
             "ACL 2024 | arXiv:2307.07924 | Status: Peer-Reviewed",
             "app/core/state_machine.py",
             "Simulates software development phases (designing, coding, testing, documenting) through multi-agent chat chains and structured state machines.")
        ]),

        ("Pillar 7: Vision, Screen Perception & OCR Grounding", [
            ("24. OmniParser: A Unified Screen Parsing Framework for Pure Vision Based GUI Agents",
             "Lu, Y., Yang, J., Shen, Y., & Awadallah, A. (2024)",
             "arXiv:2408.00203 (Microsoft Research) | Status: Preprint",
             "app/vision/vision_engine.py",
             "Converts raw screenshots into structured interactable regions and functional descriptions via YOLOv8 and BLIP-2, providing robust bounding boxes for UI grounding."),

            ("25. ScreenAgent: A Vision Language Model-driven Computer Control Agent",
             "Niu, R., Li, J., Wang, Q., Bi, Y., Kong, X., et al. (2024)",
             "IJCAI 2024 | arXiv:2402.07945 | Status: Peer-Reviewed",
             "app/vision/vision_analyzer.py",
             "Implements a vision-language perception-action loop with an experiential reflection phase to confirm visual execution success across desktop operating systems."),

            ("26. SeeClick: Harnessing GUI Grounding for Advanced Visual GUI Agents",
             "Cheng, K., Sun, Q., Chu, Y., He, H., Sun, X., et al. (2024)",
             "ACL 2024 | arXiv:2401.10935 | Status: Peer-Reviewed",
             "app/vision/vision_engine.py, app/tools/window_tools.py",
             "Trains vision-language models to directly predict normalized (x, y) click coordinates from natural language commands without accessibility trees."),

            ("27. An Overview of the Tesseract OCR Engine",
             "Smith, R. (2007)",
             "ICDAR 2007 | DOI: 10.1109/ICDAR.2007.4376991 | Status: Peer-Reviewed",
             "app/vision/ocr_reader.py, app/vision/error_reader.py",
             "Details the multi-stage optical character recognition pipeline used by Tesseract. Direct technical foundation for KLAUSE's local OCR pipeline to extract terminal stack traces and VS Code errors.")
        ]),

        ("Pillar 8: Web Navigation, Voice Synthesis & Embodied Interaction", [
            ("28. WebArena: A Realistic Web Environment for Building Autonomous Agents",
             "Zhou, S., Xu, F. F., Zhu, H., Zhou, X., Lo, R., et al. (2024)",
             "ICLR 2024 | arXiv:2307.13854 | Status: Peer-Reviewed",
             "app/agents/browser_agent.py, app/tools/browser_tools.py",
             "End-to-end benchmark for web agents interacting via Playwright and Chrome DevTools Protocol, emphasizing accessibility tree parsing and URL navigation."),

            ("29. Robust Speech Recognition via Large-Scale Weak Supervision (Whisper)",
             "Radford, A., Kim, J. W., Xu, T., Brockman, G., McLeavey, C., & Sutskever, I. (2023)",
             "ICML 2023 | arXiv:2212.04356 | Status: Peer-Reviewed",
             "app/voice/stt.py",
             "Foundational speech Transformer model behind KLAUSE's push-to-talk voice pipeline, implemented locally via faster-whisper CTranslate2 quantization."),

            ("30. Personality Traits in Large Language Models",
             "Safdari, M., Serapio-García, G., Crepy, C., Fitz, S., Romero, P., Sun, L., et al. (2023)",
             "arXiv:2307.00184 (Google DeepMind) | Status: Preprint",
             "app/config/disposition.py, disposition.yaml",
             "Empirically proves that Big Five personality traits can be reliably shaped and steered in LLMs via trait prompting, providing the basis for KLAUSE's adaptive persona engine.")
        ])
    ]

    for pillar_title, papers in pillars:
        add_header_styled(doc, pillar_title, level=2)
        for title, authors, venue, mapping, takeaway in papers:
            add_header_styled(doc, title, level=3)
            p_auth = doc.add_paragraph()
            r_auth = p_auth.add_run(f"Authors: {authors}\nVenue / Status: {venue}\nKLAUSE Subsystem: {mapping}")
            r_auth.font.size = Pt(9.5)
            r_auth.font.color.rgb = RGBColor(70, 70, 70)
            
            p_takeaway = doc.add_paragraph()
            r_label = p_takeaway.add_run("Summary & Methodological Takeaway: ")
            r_label.bold = True
            p_takeaway.add_run(takeaway)

    # 3. Comparative Gap Analysis Table
    add_header_styled(doc, "3. Comparative Gap Analysis Table", level=1)
    doc.add_paragraph(
        "The following matrix maps each of the 30 papers to its specific research or practical gap, "
        "and details the concrete architectural mechanism through which KLAUSE resolves it:"
    )

    table_data = [
        ("1", "ReAct (Yao et al., 2023)", "Agentic Loop", 
         "Tested in toy text environments; lacks integration with long-term OS state or real-time event loops.",
         "Implements ReAct inside an asynchronous FastAPI + WebSocket event dispatcher coordinating live OS events, terminal I/O, and user interrupts."),
        
        ("2", "Reflexion (Shinn et al., 2023)", "Self-Correction", 
         "Reflections remain in ephemeral prompt buffers and vanish upon process termination.",
         "Persists post-mortem reflections and error patterns permanently in MongoDB and vector storage for cross-session recall."),
        
        ("3", "Tree of Thoughts (Yao et al., 2023)", "Search & Planning", 
         "Branching factor causes combinatorial explosion, making it too slow and costly for live desktop use.",
         "Replaces unconstrained tree search with structured build phases (KLAUSE_Build_Plan.md) and pre-scoped tool categories."),
        
        ("4", "CoALA (Sumers et al., 2024)", "Cognitive Architecture", 
         "Conceptual taxonomy without a concrete production implementation for developer workstations.",
         "Directly implements CoALA's cognitive model as a functional Windows application with SQLite/Mongo, ChromaDB, and real OS toolsets."),
        
        ("5", "Toolformer (Schick et al., 2023)", "Tool Calling", 
         "Requires model fine-tuning for fixed APIs; cannot adapt to user-installed programs or changing CLIs.",
         "Employs dynamic, schema-driven Pydantic tool definitions (app/tools/base.py) requiring zero model fine-tuning."),
        
        ("6", "Gorilla (Patil et al., 2023)", "API Retrieval", 
         "Focuses on cloud REST APIs; fails on desktop window handles, subprocess stdout streams, or OS permissions.",
         "Wraps native Windows APIs (pywin32, pygetwindow) with argument sanitization specifically tuned for local binaries."),
        
        ("7", "ToolLLM (Qin et al., 2024)", "Tool Benchmarking", 
         "Evaluated on synthetic RapidAPI endpoints; fragile under network drops or desktop latency.",
         "Implements local fallback execution and a self-contained skill repository (skill_repo.py) running on the local machine."),
        
        ("8", "MemGPT (Packer et al., 2023)", "Hierarchical Memory", 
         "Aimed at conversational chatbot personas; lacks native integration with developer knowledge vaults or IDE codebases.",
         "Adapts virtual memory paging to bidirectional markdown vaults (Obsidian) and live git repository contexts."),
        
        ("9", "RAG (Lewis et al., 2020)", "Semantic Search", 
         "Treats documents as flat text chunks, losing directory structures, file hierarchy, and code symbol relations.",
         "Augments ChromaDB vector search with path-aware metadata, YAML frontmatter filtering, and per-project namespaces."),
        
        ("10", "Generative Agents (Park et al., 2023)", "Memory Stream", 
         "Simulated in an RPG sandbox; scoring heuristics are not tuned for code syntax or engineering events.",
         "Tailors the memory stream to developer events: git commits, active window switches, terminal failures, and code review notes."),
        
        ("11", "Self-RAG (Asai et al., 2024)", "Adaptive Retrieval", 
         "Requires custom pre-trained reflection tokens, making it incompatible with proprietary commercial LLM APIs.",
         "Implements an algorithmic verification layer in knowledge_base.py evaluating retrieved chunk relevance before prompt ingestion."),
        
        ("12", "GraphRAG (Edge et al., 2024)", "Graph Summarization", 
         "High LLM cost and latency to compute hierarchical communities across massive corpora; outputs static text reports.",
         "Leverages user-created Obsidian wikilinks and #tags to bootstrap graphs instantly without full-corpus LLM passes, outputting interactive .canvas maps."),
        
        ("13", "HippoRAG (Gutiérrez et al., 2024)", "Associative Memory", 
         "Graph extraction requires complex offline processing; does not dynamically update when user files change on disk.",
         "Features live file-watcher synchronization (obsidian_connector.py) incrementally updating note nodes and edge mappings upon markdown saves."),
        
        ("14", "MindMap (Wen et al., 2024)", "Graph Prompting", 
         "Theoretical study on benchmark QA; lacks real-time UI tools for users to inspect or manually edit the knowledge graph.",
         "Directly integrates with Obsidian's native visual canvas format, allowing developers to visually inspect, reorganize, and manually link agent reasoning nodes."),
        
        ("15", "KG-Agent (Jiang et al., 2025)", "KG Reasoning", 
         "Assumes structured RDF/Knowledge Bases; unable to parse messy, semi-structured personal developer notes.",
         "Incorporates robust frontmatter regex parsing, markdown AST extraction, and inline tag extractors capable of ingesting arbitrary unstructured notes."),
        
        ("16", "UFO (Zhang et al., 2024)", "Windows OS Control", 
         "Relies on cloud vision APIs for every click, leading to high latency and heavy API token costs.",
         "Implements a hybrid automation pipeline: prioritizes deterministic window APIs and CLI scripts first, falling back to vision/OCR only when necessary."),
        
        ("17", "Windows Agent Arena (Bonatti et al., 2024)", "OS Benchmarking", 
         "Benchmark environment only; does not provide a persistent personal assistant that learns user habits over time.",
         "Acts as an ambient, persistent desktop companion with long-term memory, custom user rules (rules.md), and automated event scheduling."),
        
        ("18", "OSWorld (Xie et al., 2024)", "OS Multi-Agent", 
         "Exposes high failure rates (~88%) of vision-only agents on desktop UI interactions.",
         "Bypasses brittle visual clicks by prioritizing direct OS-level execution (terminal_tools.py, system_tools.py, vscode_tools.py)."),
        
        ("19", "Cradle (Tan et al., 2024)", "Computer Control", 
         "Explores gaming environments (RDR2); lacks developer tool integration, IDE awareness, or code-syntax understanding.",
         "Tailors computer control specifically to developer workflows: managing VS Code windows, reading terminal errors, and switching active windows."),
        
        ("20", "SWE-bench (Jimenez et al., 2024)", "Autonomous SE", 
         "Evaluates offline patch generation; does not support real-time user-in-the-loop pair programming or voice steering.",
         "Provides interactive streaming chat, push-to-talk voice interrupts, and human-in-the-loop confirmation before applying destructive git commands."),
        
        ("21", "SWE-agent (Yang et al., 2024)", "Agent-Computer Interface", 
         "ACI operates exclusively inside headless Docker containers; cannot interact with native host desktop IDEs or GUI tools.",
         "Bridges the ACI concept directly into the user’s live Windows desktop, reading active VS Code windows and local terminal processes."),
        
        ("22", "MetaGPT (Hong et al., 2024)", "SOP Engineering", 
         "Multi-role SOPs generate hundreds of messages, creating excessive latency unsuited for instant desktop assistance.",
         "Condenses SOPs into an instantaneous local rule engine (app/automation/rule_engine.py) and declarative disposition settings (disposition.yaml)."),
        
        ("23", "ChatDev (Qian et al., 2024)", "Virtual Software Dev", 
         "Rigid waterfall stage progression; cannot easily handle spontaneous ad-hoc developer questions or quick file edits.",
         "Uses an adaptive state machine (app/core/state_machine.py) that switches between rapid single-turn execution and multi-step build plans."),
        
        ("24", "OmniParser (Lu et al., 2024)", "Screen Parsing", 
         "Heavy visual inference overhead (running YOLOv8 + BLIP-2 simultaneously on high-resolution multi-monitor desktops).",
         "Combines lightweight window boundary clipping (pygetwindow) with local Tesseract OCR, triggering heavy visual parsers only for unrecognized UI regions."),
        
        ("25", "ScreenAgent (Niu et al., 2024)", "Visual Control", 
         "Visual reflection loop is prone to infinite loops when UI elements do not change visually after failed clicks.",
         "Employs explicit timeout thresholds, process exit code checks, and fallback CLI queries when visual confirmation fails."),
        
        ("26", "SeeClick (Cheng et al., 2024)", "GUI Grounding", 
         "Sub-pixel coordinate prediction frequently drifts on high-DPI Windows displays or multi-monitor configurations.",
         "Implements DPI-aware coordinate scaling and OS-level window focusing before dispatching simulated mouse/keyboard events."),
        
        ("27", "Tesseract OCR (Smith, 2007)", "OCR Processing", 
         "Traditional OCR engine with zero semantic understanding of modern IDE syntax themes, stack traces, or terminal ANSI codes.",
         "Enhances raw Tesseract output with regex filtering (ERROR_PATTERNS in error_reader.py) and passes cleaned text to LLMs for diagnostic interpretation."),
        
        ("28", "WebArena (Zhou et al., 2024)", "Web Automation", 
         "Benchmarks web tasks in isolated sandboxes; disconnected from local user credentials, existing browser sessions, or local files.",
         "Connects directly to existing user browsers (Brave, Chrome, Edge) via Chrome DevTools Protocol (CDP_PORT = 9222), preserving sessions and cookies."),
        
        ("29", "Whisper (Radford et al., 2023)", "Speech-to-Text", 
         "Standard PyTorch models introduce noticeable latency on desktop hardware, interfering with fluid conversational flow.",
         "Integrates faster-whisper (CTranslate2 INT8 quantization) on local threads, delivering low-latency push-to-talk transcription."),
        
        ("30", "Personality Traits (Safdari et al., 2023)", "Synthetic Personality", 
         "Studies static prompt priming on benchmark datasets; lacks dynamic personality adaptation based on user workflow state.",
         "Implements a real-time disposition controller (app/config/disposition.py) that modulates assistant tone based on task urgency and error states.")
    ]

    # Create table
    table = doc.add_table(rows=len(table_data) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths in inches
    col_widths = [Inches(0.4), Inches(1.5), Inches(1.1), Inches(2.3), Inches(2.7)]

    # Header row
    headers = ["#", "Paper Citation", "Core Domain", "Identified Research / Practical Gap", "How KLAUSE Resolves the Gap"]
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "182B49") # Navy
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row_vals in enumerate(table_data):
        row = table.rows[r_idx + 1]
        bg_color = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx in (0, 1):
                r.font.bold = True

    # Section 4: Architectural Conclusions
    add_header_styled(doc, "4. Architectural Synthesis & Conclusions", level=1)
    doc.add_paragraph(
        "By systematically evaluating these 30 foundational contributions, three clear engineering insights emerge that validate KLAUSE's design:\n"
        "1. Hybrid Execution Beats Pure Vision: Pure vision-based OS control (as observed in OSWorld and Cradle) exhibits unacceptably high error rates on desktop UI controls. KLAUSE achieves resilience by prioritizing deterministic Windows APIs, CLI tools, and Playwright CDP before falling back to OCR and visual grounding.\n"
        "2. Hierarchical Local Memory as a First-Class Citizen: Moving beyond cloud-centric agent models, KLAUSE executes MemGPT-style virtual memory paging grounded in the developer's local filesystem (MongoDB, ChromaDB, and Obsidian Vaults). This guarantees data privacy, zero vendor lock-in, and cross-session persistence.\n"
        "3. Live Desktop Grounding: Rather than trapping autonomous engineering agents in isolated cloud containers (SWE-bench / SWE-agent), KLAUSE embeds the Agent-Computer Interface directly onto the developer's workstation, integrating active IDE windows, terminal feedback loops, and local push-to-talk voice activation."
    )

    out_path = r"e:\KLAUSE\KLAUSE_Literature_Survey.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    main()
